"""Deterministic DOCX layout; pagination is delegated to the word processor."""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .models import Project, document_sections


def field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, sep, text, end):
        run._r.append(element)


def export_docx(project: Project):
    doc = Document()
    fmt = project.formatting
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left", "right", "top", "bottom"):
        setattr(section, side + "_margin", Mm(getattr(fmt, side)))
    section.footer_distance = Mm(10)
    section.different_first_page_header_footer = fmt.title_page
    pgnum = OxmlElement("w:pgNumType")
    pgnum.set(qn("w:start"), str(fmt.page_start))
    section._sectPr.append(pgnum)
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "List Bullet",
        "List Number",
        "Caption",
        "Footer",
    ):
        style = doc.styles[name]
        style.font.name = fmt.font
        style.font.size = Pt(fmt.size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.rFonts
        for attr in list(fonts.attrib):
            del fonts.attrib[attr]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), fmt.font)
        # Templates can carry borders and theme overrides even on the Title style.
        ppr = style.element.get_or_add_pPr()
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
        pf = style.paragraph_format
        pf.space_after = Pt(0)
        pf.line_spacing = fmt.spacing
    normal = doc.styles["Normal"].paragraph_format
    normal.first_line_indent = Mm(fmt.indent)
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Heading 1", "Heading 2"):
        pf = doc.styles[name].paragraph_format
        pf.first_line_indent = Mm(0)
        pf.space_before = Pt(12)
        pf.space_after = Pt(8)
        pf.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Mm(0)
    if fmt.page_numbers:
        field(footer, " PAGE ", str(fmt.page_start))
    if fmt.title_page:

        def title_line(text, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, style=None):
            p = doc.add_paragraph(text, style)
            p.alignment = align
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.space_after = Pt(after)
            return p

        title_line(project.institution or "Образовательная организация", 90)
        title_line("ПРАКТИЧЕСКАЯ РАБОТА", 16, style="Title")
        title_line(project.topic or project.title, 12)
        title_line(("Дисциплина: " + project.subject) if project.subject else "", 48)
        for label, value in [
            ("Выполнил(а)", project.student),
            ("Группа", project.group),
            ("Проверил(а)", project.teacher),
        ]:
            if value:
                title_line(f"{label}: {value}", 8, WD_ALIGN_PARAGRAPH.RIGHT)
        title_line(" ".join(filter(None, [project.city, project.year])), 0)
        doc.add_page_break()
    if fmt.toc:
        heading = doc.add_paragraph("Содержание")
        heading.paragraph_format.first_line_indent = Mm(0)
        heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.keep_with_next = True
        heading.runs[0].bold = True
        p = doc.add_paragraph()
        field(p, ' TOC \\o "1-2" \\h \\z \\u ', "Обновите содержание в Word: Ctrl+A, затем F9")
        doc.add_page_break()
    table_number = 0
    for index, item in enumerate(document_sections(project)):
        if item.new_page and index:
            doc.add_page_break()
        doc.add_heading(item.title, 1)
        for block in item.blocks:
            if block.type == "paragraph":
                doc.add_paragraph(block.text)
            elif block.type in ("list", "number_list"):
                for text in block.items:
                    p = doc.add_paragraph(
                        text, "List Number" if block.type == "number_list" else "List Bullet"
                    )
                    p.paragraph_format.first_line_indent = Mm(0)
            elif block.type == "page_break":
                doc.add_page_break()
            elif block.type == "table":
                table_number += 1
                p = doc.add_paragraph(f"Таблица {table_number} — {block.caption or 'Данные'}")
                p.paragraph_format.first_line_indent = Mm(0)
                p.paragraph_format.keep_with_next = True
                table = doc.add_table(rows=1, cols=len(block.headers))
                table.style = "Table Grid"
                width = 210 - fmt.left - fmt.right
                table.autofit = False
                for col in table.columns:
                    col.width = Mm(width / len(block.headers))
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for i, value in enumerate(block.headers):
                    table.rows[0].cells[i].text = value
                for row in block.rows:
                    cells = table.add_row().cells
                    for i, value in enumerate(row):
                        cells[i].text = value
                for ri, row in enumerate(table.rows):
                    cant_split = OxmlElement("w:cantSplit")
                    row._tr.get_or_add_trPr().append(cant_split)
                    for cell in row.cells:
                        cell.width = Mm(width / len(block.headers))
                        for p in cell.paragraphs:
                            p.paragraph_format.first_line_indent = Mm(0)
                            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            p.paragraph_format.line_spacing = 1
                            for run in p.runs:
                                run.font.size = Pt(min(fmt.size, 12))
                                run.bold = ri == 0
    if project.sources.strip():
        doc.add_heading("Список использованных источников", 1)
        for line in project.sources.splitlines():
            if line.strip():
                p = doc.add_paragraph(line.strip(), "List Number")
                p.paragraph_format.first_line_indent = Mm(0)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    result = BytesIO()
    doc.save(result)
    return result.getvalue()
