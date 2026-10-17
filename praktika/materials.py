"""Text extraction from user attachments, independent of HTTP transport."""

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .models import Attachment


def extract_attachment(filename: str, raw: bytes) -> Attachment:
    name = Path(filename or "Материал").name
    ext = Path(name).suffix.lower()
    try:
        if ext == ".docx":
            doc = Document(BytesIO(raw))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    " | ".join(c.text for c in row.cells)
                    for table in doc.tables
                    for row in table.rows
                ]
            )
        elif ext == ".pdf":
            pdf = PdfReader(BytesIO(raw))
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif ext in (".txt", ".md"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1251")
        else:
            raise ValueError("Поддерживаются PDF, DOCX, TXT и Markdown")
    except Exception as error:
        if isinstance(error, ValueError):
            raise
        raise ValueError(
            "Не удалось прочитать файл. Проверьте формат, пароль и целостность документа."
        )
    if not text.strip():
        raise ValueError("В файле нет извлекаемого текста. Для сканов сначала выполните OCR.")
    if len(text) > 1_000_000:
        raise ValueError("Слишком много текста: разделите файл на части (до 1 млн символов)")
    return Attachment(name=name, text=text)
