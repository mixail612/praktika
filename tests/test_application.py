import io
import os
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

TEST_DATA = tempfile.TemporaryDirectory()
os.environ["PRAKTIKA_DATA"] = TEST_DATA.name
from docx import Document
from fastapi.testclient import TestClient

from praktika import generation, storage
from praktika.documents import export_docx
from praktika.models import Block, Formatting, Profile, Project, Section
from praktika.server import app


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def wait_job(self, identifier):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            job = self.client.get(f"/api/projects/{identifier}/job").json()
            if job and job["status"] not in ("running", "cancelling"):
                return job
            time.sleep(0.02)
        self.fail("Generation did not finish")

    def test_global_router_and_last_model(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(storage, "DATA", Path(folder)):
            first = storage.write(
                "profiles", Profile(name="First", base_url="https://example.com/v1")
            )
            second = storage.write(
                "profiles", Profile(name="Second", base_url="https://example.org/v1")
            )
            self.assertEqual(
                self.client.put("/api/settings", json={"profile_id": first.id}).status_code, 200
            )
            work = self.client.post("/api/projects").json()
            self.assertEqual(work["profile_id"], first.id)
            work["model"] = "test-model"
            self.client.put("/api/projects/" + work["id"], json=work).raise_for_status()
            self.assertEqual(self.client.post("/api/projects").json()["model"], "test-model")
            self.client.put("/api/settings", json={"profile_id": second.id}).raise_for_status()
            saved = self.client.put("/api/projects/" + work["id"], json=work).json()
            self.assertEqual(saved["profile_id"], second.id)
            self.assertEqual(self.client.get("/api/settings").json()["profile_id"], second.id)

    def test_project_roundtrip_export_import(self):
        p = self.client.post("/api/projects").json()
        p["topic"] = "Алгоритмы сортировки"
        p["sections"] = [
            Section(
                title="Сравнение",
                blocks=[
                    Block(text="Текст работы"),
                    Block(
                        type="table",
                        caption="Сравнение",
                        headers=["Метод", "Сложность"],
                        rows=[["Сортировка", "O(n)"]],
                    ),
                    Block(type="list", items=["Первый пункт"]),
                ],
            ).model_dump()
        ]
        response = self.client.put("/api/projects/" + p["id"], json=p)
        self.assertEqual(response.status_code, 200)
        response = self.client.get("/api/projects/" + p["id"] + "/export.docx")
        self.assertEqual(response.status_code, 200)
        doc = Document(io.BytesIO(response.content))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "O(n)")
        self.assertAlmostEqual(doc.sections[0].left_margin.mm, 30, places=1)
        raw = self.client.get("/api/projects/" + p["id"] + "/export.json").content
        imported = self.client.post(
            "/api/import", files={"file": ("backup.json", raw, "application/json")}
        )
        self.assertEqual(imported.status_code, 200)
        self.assertNotEqual(imported.json()["id"], p["id"])
        self.assertEqual(imported.json()["topic"], p["topic"])

    def test_legacy_import_and_extraction(self):
        raw = (Path(__file__).parent / "fixtures" / "legacy_work.json").read_bytes()
        response = self.client.post("/api/import", files={"file": ("old.json", raw)})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["sections"])
        response = self.client.post(
            "/api/extract", files={"file": ("context.txt", "Русский текст".encode())}
        )
        self.assertEqual(response.json()["text"], "Русский текст")
        doc = export_docx(
            Project(sections=[Section(title="Тест", blocks=[Block(text="Материал для модели")])])
        )
        response = self.client.post("/api/extract", files={"file": ("context.docx", doc)})
        self.assertIn("Материал для модели", response.json()["text"])

    def test_validation_and_local_origin(self):
        self.assertEqual(
            self.client.post(
                "/api/projects", headers={"Origin": "https://foreign.example"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.get("/api/health", headers={"Host": "foreign.example"}).status_code, 400
        )
        response = self.client.post("/api/profiles", json={"base_url": "not a url"})
        self.assertEqual(response.status_code, 400)
        p = self.client.post("/api/projects").json()
        p["pages"] = 0
        self.assertEqual(self.client.put("/api/projects/" + p["id"], json=p).status_code, 422)
        self.assertEqual(
            self.client.post("/api/extract", files={"file": ("script.exe", b"data")}).status_code,
            400,
        )

    def test_document_composition_switches(self):
        from zipfile import ZipFile

        project = Project(
            formatting=Formatting(
                title_page=False, toc=False, page_numbers=False, conclusion=False
            ),
            sections=[
                Section(title="Работа", blocks=[Block(text="Основной текст")]),
                Section(title="Вывод", blocks=[Block(text="Итоговый текст")]),
            ],
        )
        raw = export_docx(project)
        with ZipFile(io.BytesIO(raw)) as z:
            xml = z.read("word/document.xml").decode()
            footer = z.read("word/footer1.xml").decode()
        self.assertNotIn("Итоговый текст", xml)
        self.assertNotIn("ПРАКТИЧЕСКАЯ РАБОТА", xml)
        self.assertNotIn(" TOC ", xml)
        self.assertNotIn(" PAGE ", footer)
        project.formatting = Formatting(
            title_page=True, toc=True, page_numbers=True, conclusion=True
        )
        with ZipFile(io.BytesIO(export_docx(project))) as z:
            xml = z.read("word/document.xml").decode()
            footer = z.read("word/footer1.xml").decode()
        self.assertIn("Итоговый текст", xml)
        self.assertIn("ПРАКТИЧЕСКАЯ РАБОТА", xml)
        self.assertIn(" TOC ", xml)
        self.assertIn(" PAGE ", footer)

    def test_key_reveal_is_explicit_and_does_not_modify_storage(self):
        profile = storage.write("profiles", Profile())
        with (
            patch.object(storage, "get_key", return_value="saved-test-secret"),
            patch.object(storage, "save_key") as save,
        ):
            response = self.client.post(f"/api/profiles/{profile.id}/reveal-key")
            self.assertEqual(response.json()["api_key"], "saved-test-secret")
            self.assertEqual(response.headers["cache-control"], "no-store")
            self.assertNotIn("api_key", self.client.get("/api/profiles").json()[0])
            save.assert_not_called()

    def test_conclusion_added_when_requested(self):
        profile = storage.write("profiles", Profile(model="mock", allow_no_key=True))
        project = storage.write("projects", Project(topic="Тема", profile_id=profile.id))
        with (
            patch.object(generation, "client_for", return_value=unittest.mock.MagicMock()),
            patch.object(
                generation, "ask", return_value={"sections": [{"title": "Практическая часть"}]}
            ),
        ):
            self.client.post(f"/api/projects/{project.id}/generate", json={"mode": "plan"})
            self.assertEqual(self.wait_job(project.id)["status"], "complete")
        self.assertEqual(storage.read("projects", project.id).sections[-1].kind, "conclusion")

    def test_profile_secrets_are_not_returned(self):
        with (
            patch.object(storage, "save_key") as write,
            patch.object(storage, "get_key", return_value="secret-test"),
        ):
            result = self.client.post(
                "/api/profiles", json={"name": "Mock", "api_key": "secret-test"}
            ).json()
            write.assert_called_once()
            self.assertNotIn("api_key", result)
            self.assertTrue(result["has_key"])
            self.assertNotIn("secret-test", storage.path_for("profiles", result["id"]).read_text())

    def test_model_catalog_uses_draft_credentials_without_saving(self):
        from types import SimpleNamespace

        fake = unittest.mock.MagicMock()
        fake.__enter__.return_value.models.list.return_value.data = [
            SimpleNamespace(id="z/model", model_extra={"name": "Z Model"}),
            SimpleNamespace(id="a/model", model_extra={}),
            SimpleNamespace(id="a/model", model_extra={}),
        ]
        profile = Profile()
        with patch.object(generation, "client_for", return_value=fake) as create:
            response = self.client.post(
                "/api/models", json={**profile.model_dump(), "api_key": "draft-secret"}
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual([m["id"] for m in response.json()["models"]], ["a/model", "z/model"])
        self.assertNotIn("draft-secret", response.text)
        self.assertFalse(storage.path_for("profiles", profile.id).exists())
        self.assertEqual(create.call_args.kwargs["api_key"], "draft-secret")

    def test_router_requires_key_before_network_and_saves_catalog(self):
        from types import SimpleNamespace

        with (
            patch.object(storage, "get_key", return_value=""),
            patch.object(generation, "client_for") as factory,
        ):
            response = self.client.post(
                "/api/routers/connect", json={"base_url": "https://router.example/v1"}
            )
            self.assertEqual(response.status_code, 400)
            factory.assert_not_called()
        fake = unittest.mock.MagicMock()
        fake.__enter__.return_value.models.list.return_value.data = [
            SimpleNamespace(id="router/model", model_extra={})
        ]
        with (
            patch.object(generation, "client_for", return_value=fake),
            patch.object(storage, "save_key"),
            patch.object(storage, "get_key", return_value="key"),
        ):
            response = self.client.post(
                "/api/routers/connect",
                json={"base_url": "https://router.example/v1", "api_key": "key"},
            )
        self.assertEqual(response.status_code, 200)
        result = response.json()["profile"]
        self.assertEqual(result["model"], "")
        self.assertEqual(result["models"][0]["id"], "router/model")
        self.assertTrue(result["catalog_supported"])
        self.assertNotIn("api_key", result)

    def test_catalog_uses_saved_key_in_authorization_header(self):
        import httpx
        from openai import OpenAI

        observed = []

        def handler(request):
            observed.append(request.headers.get("authorization"))
            return httpx.Response(
                200,
                json={
                    "object": "list",
                    "data": [
                        {"id": "model-a", "object": "model", "created": 0, "owned_by": "router"}
                    ],
                },
            )

        real_client = OpenAI(
            api_key="saved-router-key",
            base_url="https://router.example/v1",
            http_client=httpx.Client(transport=httpx.MockTransport(handler)),
        )
        with (
            patch.object(storage, "get_key", return_value="saved-router-key"),
            patch.object(generation, "client_for", return_value=real_client),
        ):
            result = self.client.post(
                "/api/models", json={"base_url": "https://router.example/v1", "api_key": ""}
            )
        self.assertEqual(result.status_code, 200)
        self.assertEqual(observed, ["Bearer saved-router-key"])

    def test_reasoning_and_model_override_are_sent(self):
        from types import SimpleNamespace

        fake = unittest.mock.MagicMock()
        fake.chat.completions.create.return_value.choices = [
            SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content='{"blocks":[]}'))
        ]
        generation.ask(fake, Profile(model="test-model", reasoning_effort="high"), [])
        self.assertEqual(fake.chat.completions.create.call_args.kwargs["reasoning_effort"], "high")
        generation.ask(fake, Profile(), [])
        self.assertNotIn("reasoning_effort", fake.chat.completions.create.call_args.kwargs)
        profile = storage.write(
            "profiles", Profile(model="default-model", reasoning_effort="low", allow_no_key=True)
        )
        p = storage.write(
            "projects",
            Project(
                formatting=Formatting(conclusion=False),
                topic="Тест",
                profile_id=profile.id,
                model="chosen-model",
                reasoning_effort="high",
            ),
        )
        with patch.object(generation.POOL, "submit") as submit:
            job = generation.start(p.id, "plan")
            effective = submit.call_args.args[3]
            self.assertEqual(effective.model, "chosen-model")
            self.assertEqual(effective.reasoning_effort, "high")
            generation.JOBS[job["id"]]["status"] = "complete"

    def test_generation_plan_and_sections(self):
        profile = storage.write("profiles", Profile(name="Mock", model="mock", allow_no_key=True))
        p = storage.write(
            "projects",
            Project(formatting=Formatting(conclusion=False), topic="Тест", profile_id=profile.id),
        )
        fake_client = unittest.mock.MagicMock()
        responses = [
            {"sections": [{"title": "Практическая часть", "description": "Проверить алгоритм"}]},
            {"blocks": [{"type": "paragraph", "text": "Результат генерации"}]},
        ]
        with (
            patch.object(generation, "client_for", return_value=fake_client),
            patch.object(generation, "ask", side_effect=responses),
        ):
            self.assertEqual(
                self.client.post(
                    f"/api/projects/{p.id}/generate", json={"mode": "plan"}
                ).status_code,
                200,
            )
            self.assertEqual(self.wait_job(p.id)["status"], "complete")
            self.assertEqual(
                self.client.post(
                    f"/api/projects/{p.id}/generate", json={"mode": "full"}
                ).status_code,
                200,
            )
            self.assertEqual(self.wait_job(p.id)["status"], "complete")
        result = storage.read("projects", p.id)
        self.assertEqual(result.sections[0].blocks[0].text, "Результат генерации")

    def test_cancel_discards_response_and_blocks_editing(self):
        profile = storage.write("profiles", Profile(name="Mock", model="mock", allow_no_key=True))
        p = storage.write(
            "projects",
            Project(
                formatting=Formatting(conclusion=False),
                topic="Тест",
                profile_id=profile.id,
                sections=[Section(title="Раздел", blocks=[Block(text="Исходный текст")])],
            ),
        )
        started = threading.Event()
        release = threading.Event()

        def fake_ask(*args):
            started.set()
            release.wait(3)
            return {"blocks": [{"type": "paragraph", "text": "Отброшенный ответ"}]}

        with (
            patch.object(generation, "client_for", return_value=unittest.mock.MagicMock()),
            patch.object(generation, "ask", side_effect=fake_ask),
        ):
            result = self.client.post(
                f"/api/projects/{p.id}/generate", json={"mode": "full"}
            ).json()
            self.assertTrue(started.wait(2))
            self.assertEqual(
                self.client.put(f"/api/projects/{p.id}", json=p.model_dump()).status_code, 409
            )
            self.client.post("/api/jobs/" + result["id"] + "/cancel")
            release.set()
            self.assertEqual(self.wait_job(p.id)["status"], "cancelled")
        self.assertEqual(
            storage.read("projects", p.id).sections[0].blocks[0].text, "Исходный текст"
        )

    def test_invalid_model_response_preserves_previous_content(self):
        profile = storage.write("profiles", Profile(name="Mock", model="mock", allow_no_key=True))
        p = storage.write(
            "projects",
            Project(
                formatting=Formatting(conclusion=False),
                topic="Тест",
                profile_id=profile.id,
                sections=[Section(title="Раздел", blocks=[Block(text="Сохранённый текст")])],
            ),
        )
        with (
            patch.object(generation, "client_for", return_value=unittest.mock.MagicMock()),
            patch.object(
                generation,
                "ask",
                return_value={
                    "blocks": [{"type": "table", "headers": ["A", "B"], "rows": [["invalid"]]}]
                },
            ),
        ):
            self.client.post(f"/api/projects/{p.id}/generate", json={"mode": "full"})
            self.assertEqual(self.wait_job(p.id)["status"], "failed")
        self.assertEqual(
            storage.read("projects", p.id).sections[0].blocks[0].text, "Сохранённый текст"
        )


if __name__ == "__main__":
    unittest.main()
